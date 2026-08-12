"""環境省の規制対象種一覧 (xlsx) から、同梱する種リストを作り直す。

実行時に env.go.jp を叩かないため、リストはこのリポジトリに同梱する。告示は年単位で
しか変わらないので、更新もこの手作業でよい。

    uv run --with openpyxl python data/build_protected_species.py

対象は別表第一 (国内希少野生動植物種) と緊急指定種の2枚だけ。別表第二 (国際希少
野生動植物種) は入れない。あちらが規制するのは主に譲渡し・輸出入で、日本の野外で
出会う生きものの話ではないため、写真から名前を調べる画面で出すと関係のない警告になる。
"""

from __future__ import annotations

import io
import json
import re
from datetime import date
from pathlib import Path
from urllib.request import urlopen

SOURCE_PAGE = "https://www.env.go.jp/nature/kisho/domestic/list.html"
WORKBOOK_URL = "https://www.env.go.jp/content/000046164.xlsx"
OUTPUT = Path(__file__).parent / "protected_species.json"

# 指定なしの欄には「－」と「ー」が混在している (全角ハイフンと長音符)。
# 指定ありの印だけを見れば、この揺れを踏まない。
DESIGNATED = "○"


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _subtype(first: object, second: object) -> str:
    if DESIGNATED in _clean(first):
        return "特定第一種"
    if DESIGNATED in _clean(second):
        return "特定第二種"
    return ""


def build() -> dict:
    import openpyxl

    raw = urlopen(WORKBOOK_URL, timeout=120).read()  # noqa: S310 - 環境省の固定URL
    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)

    domestic_sheet = next(s for s in workbook.sheetnames if s.startswith("国内希少野生動植物種"))
    urgent_sheet = next(s for s in workbook.sheetnames if s.startswith("緊急指定種"))

    species: list[dict] = []
    seen: set[str] = set()

    for row in workbook[domestic_sheet].iter_rows(min_row=3, values_only=True):
        name = _clean(row[10])
        if not name or name.casefold() in seen:
            continue
        seen.add(name.casefold())
        species.append({
            "scientific_name": name,
            "common_name_ja": _clean(row[11]),
            "category": "国内希少野生動植物種",
            "subtype": _subtype(row[13], row[14]),
        })

    for row in workbook[urgent_sheet].iter_rows(min_row=3, values_only=True):
        name = _clean(row[2])
        if not name or name.casefold() in seen:
            continue
        seen.add(name.casefold())
        species.append({
            "scientific_name": name,
            "common_name_ja": _clean(row[3]),
            "category": "緊急指定種",
            "subtype": "",
        })

    return {
        "source": "環境省「種の保存法における規制対象種一覧」",
        "source_page": SOURCE_PAGE,
        "workbook_url": WORKBOOK_URL,
        "sheets": [domestic_sheet, urgent_sheet],
        "as_of": "令和7年2月",
        "retrieved_on": date.today().isoformat(),
        "count": len(species),
        "species": sorted(species, key=lambda item: item["scientific_name"]),
    }


if __name__ == "__main__":
    data = build()
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{OUTPUT}: {data['count']}件 ({', '.join(data['sheets'])})")
